"""
Web scraper: book catalogue -> formatted Excel file.

Scrapes https://books.toscrape.com, a public site built for scraping practice.
Follows pagination automatically and collects, for every book:
title, price, star rating, stock status and product URL.

Usage:
    python scraper.py                 # all 50 pages (1,000 books)
    python scraper.py --pages 3       # first 3 pages only (quick test)
    python scraper.py --output my.xlsx
"""

import argparse
import time
from pathlib import Path
from urllib.parse import urljoin

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://books.toscrape.com/catalogue/page-1.html"
HEADERS = {"User-Agent": "Mozilla/5.0 (portfolio scraper; polite, 1 req/sec)"}
RATING_WORDS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}


def fetch(session: requests.Session, url: str, retries: int = 3) -> str:
    """Download a page, retrying a few times on network errors."""
    for attempt in range(1, retries + 1):
        try:
            resp = session.get(url, headers=HEADERS, timeout=15)
            resp.raise_for_status()
            resp.encoding = "utf-8"
            return resp.text
        except requests.RequestException as err:
            if attempt == retries:
                raise
            print(f"  retry {attempt} for {url}: {err}")
            time.sleep(2 * attempt)


def parse_page(html: str, page_url: str) -> tuple[list[dict], str | None]:
    """Extract every book on one listing page, plus the next-page link."""
    soup = BeautifulSoup(html, "html.parser")
    books = []
    for card in soup.select("article.product_pod"):
        link = card.h3.a
        rating_class = card.select_one("p.star-rating")["class"]
        rating = next((RATING_WORDS[c] for c in rating_class if c in RATING_WORDS), None)
        price_text = card.select_one("p.price_color").get_text(strip=True)
        books.append({
            "Title": link["title"].strip(),
            "Price (GBP)": float(price_text.replace("£", "").replace("Â", "")),
            "Rating (1-5)": rating,
            "In Stock": "in stock" in card.select_one("p.availability").get_text().lower(),
            "URL": urljoin(page_url, link["href"]),
        })
    nxt = soup.select_one("li.next > a")
    return books, (urljoin(page_url, nxt["href"]) if nxt else None)


def scrape(max_pages: int | None = None, delay: float = 1.0) -> pd.DataFrame:
    session = requests.Session()
    url, page, rows = BASE_URL, 0, []
    while url and (max_pages is None or page < max_pages):
        page += 1
        books, url = parse_page(fetch(session, url), url)
        rows.extend(books)
        print(f"  page {page}: {len(books)} books (total {len(rows)})")
        if url:
            time.sleep(delay)  # be polite to the server
    return pd.DataFrame(rows)


def save_excel(df: pd.DataFrame, out_path: Path):
    out_path.parent.mkdir(parents=True, exist_ok=True)
    summary = (df.groupby("Rating (1-5)", as_index=False)
                 .agg(Books=("Title", "count"), Avg_Price=("Price (GBP)", "mean"))
                 .rename(columns={"Avg_Price": "Avg Price (GBP)"}))
    summary["Avg Price (GBP)"] = summary["Avg Price (GBP)"].round(2)

    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Books", index=False)
        summary.to_excel(xw, sheet_name="By Rating", index=False)
        for name, ws in xw.sheets.items():
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="1F6F43")
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center")
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 10), 55)
                if "GBP" in str(col[0].value):
                    for c in col[1:]:
                        c.number_format = '"£"#,##0.00'


def main():
    ap = argparse.ArgumentParser(description="Scrape books.toscrape.com into Excel.")
    ap.add_argument("--pages", type=int, default=None, help="limit number of pages (default: all)")
    ap.add_argument("--output", type=Path, default=Path(__file__).parent / "output" / "books_data.xlsx")
    args = ap.parse_args()

    start = time.time()
    print("Scraping books.toscrape.com ...")
    df = scrape(args.pages)
    save_excel(df, args.output)
    print(f"Done: {len(df)} books saved to {args.output} in {time.time() - start:.0f}s")


if __name__ == "__main__":
    main()
