"""
JavaScript-rendered site scraper (Playwright) -> formatted Excel.

Some websites load their content with JavaScript, so a plain HTTP request
returns an empty page. This scraper drives a real (headless) browser, waits
for the content to render, extracts it, and follows pagination.

Target: https://quotes.toscrape.com/js/  (a public site built for scraping practice)
Collects every quote's text, author and tags, then writes an Excel report with:
  - Quotes     : one row per quote
  - By Author  : number of quotes per author
  - By Tag     : how often each tag appears

Usage:
    python scraper_js.py                 # all pages
    python scraper_js.py --pages 2       # quick test
    python scraper_js.py --headed        # watch the browser work
"""

import argparse
import time
from collections import Counter
from pathlib import Path
from urllib.parse import urljoin

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.sync_api import TimeoutError as PWTimeout
from playwright.sync_api import sync_playwright

START_URL = "https://quotes.toscrape.com/js/"

# JavaScript run inside the page to read the rendered quotes
EXTRACT_JS = """
els => els.map(q => ({
    text:   (q.querySelector('.text')   || {}).innerText || '',
    author: (q.querySelector('.author') || {}).innerText || '',
    tags:   Array.from(q.querySelectorAll('.tags a.tag')).map(t => t.innerText)
}))
"""


def scrape(start_url: str, max_pages: int | None, headed: bool, delay: float = 1.0) -> list[dict]:
    rows: list[dict] = []
    with sync_playwright() as p:
        try:  # use an installed Google Chrome if there is one
            browser = p.chromium.launch(headless=not headed, channel="chrome")
        except Exception:  # otherwise Playwright's own Chromium
            browser = p.chromium.launch(headless=not headed)
        page = browser.new_page()
        url, page_no = start_url, 0
        while url and (max_pages is None or page_no < max_pages):
            page_no += 1
            for attempt in range(1, 4):  # retry slow or failed loads
                try:
                    page.goto(url, timeout=30_000)
                    page.wait_for_selector(".quote", timeout=15_000)  # wait for JS to render
                    break
                except PWTimeout:
                    if attempt == 3:
                        raise
                    print(f"  retry {attempt} for {url}")
                    time.sleep(2 * attempt)

            quotes = page.eval_on_selector_all(".quote", EXTRACT_JS)
            for q in quotes:
                rows.append({
                    "Quote": q["text"].strip().strip("“”\""),
                    "Author": q["author"].strip(),
                    "Tags": ", ".join(q["tags"]),
                    "Page": page_no,
                })
            print(f"  page {page_no}: {len(quotes)} quotes (total {len(rows)})")

            nxt = page.query_selector("li.next > a")
            url = urljoin(url, nxt.get_attribute("href")) if nxt else None
            if url:
                time.sleep(delay)  # be polite to the server
        browser.close()
    return rows


def style(ws):
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F6F43")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 10), 70)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def save_excel(rows: list[dict], out: Path):
    out.parent.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(rows)
    by_author = (df.groupby("Author").size().reset_index(name="Quotes")
                   .sort_values(["Quotes", "Author"], ascending=[False, True]))
    tag_counts = Counter(t.strip() for tags in df["Tags"] for t in tags.split(",") if t.strip())
    by_tag = (pd.DataFrame(tag_counts.items(), columns=["Tag", "Count"])
                .sort_values(["Count", "Tag"], ascending=[False, True]))
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Quotes", index=False)
        by_author.to_excel(xw, sheet_name="By Author", index=False)
        by_tag.to_excel(xw, sheet_name="By Tag", index=False)
        for ws in xw.sheets.values():
            style(ws)
    return df, by_author, by_tag


def main():
    ap = argparse.ArgumentParser(description="Scrape a JavaScript-rendered site into Excel.")
    ap.add_argument("--url", default=START_URL)
    ap.add_argument("--pages", type=int, default=None, help="limit pages (default: all)")
    ap.add_argument("--headed", action="store_true", help="show the browser window")
    ap.add_argument("--output", type=Path, default=Path(__file__).parent / "output" / "quotes_data.xlsx")
    args = ap.parse_args()

    start = time.time()
    print(f"Scraping {args.url} with a headless browser ...")
    rows = scrape(args.url, args.pages, args.headed)
    df, by_author, by_tag = save_excel(rows, args.output)
    print(f"Done: {len(df)} quotes, {len(by_author)} authors, {len(by_tag)} tags "
          f"saved to {args.output.relative_to(Path.cwd()) if args.output.is_relative_to(Path.cwd()) else args.output} "
          f"in {time.time() - start:.0f}s")


if __name__ == "__main__":
    main()
