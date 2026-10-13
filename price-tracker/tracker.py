"""
Price tracker with history, change detection and alerts -> Excel report.

Collects live prices on every run, stores them in a local SQLite database,
compares each run with the previous one, flags moves above a threshold and
rebuilds an Excel report with a price-history chart.

Data source: CoinGecko public API (free, no key needed for light use).
The same structure works for product prices scraped from any website.

Usage:
    python tracker.py                         # one run (schedule it with cron / Task Scheduler)
    python tracker.py --every 10 --runs 6     # built-in schedule: every 10 min, 6 runs
    python tracker.py --coins 15 --threshold 1.0
"""

import argparse
import sqlite3
import time
from datetime import datetime
from pathlib import Path

import pandas as pd
import requests
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
API = "https://api.coingecko.com/api/v3/coins/markets"
HEADERS = {"User-Agent": "price-tracker-portfolio/1.0", "Accept": "application/json"}

GREEN, RED = "E3EFE7", "FBEDEB"


# ---------- collect ----------

def fetch_prices(n: int) -> list[dict]:
    """Top-n coins by market cap, with retries for network hiccups and rate limits."""
    params = {"vs_currency": "usd", "order": "market_cap_desc", "per_page": n, "page": 1}
    for attempt in range(1, 4):
        try:
            r = requests.get(API, params=params, headers=HEADERS, timeout=20)
            if r.status_code == 429:  # rate limited: wait and retry
                raise requests.RequestException("rate limited (429)")
            r.raise_for_status()
            return [{"coin_id": c["id"], "symbol": c["symbol"].upper(), "name": c["name"],
                     "price": float(c["current_price"]), "market_cap": float(c["market_cap"] or 0)}
                    for c in r.json()]
        except requests.RequestException as err:
            if attempt == 3:
                raise
            print(f"  retry {attempt}: {err}")
            time.sleep(10 * attempt)


# ---------- store ----------

def connect(db: Path) -> sqlite3.Connection:
    con = sqlite3.connect(db)
    con.execute("""CREATE TABLE IF NOT EXISTS snapshots (
                       run_at TEXT, coin_id TEXT, symbol TEXT, name TEXT,
                       price REAL, market_cap REAL)""")
    return con


def save_snapshot(con, rows: list[dict], run_at: str):
    con.executemany("INSERT INTO snapshots VALUES (?,?,?,?,?,?)",
                    [(run_at, r["coin_id"], r["symbol"], r["name"], r["price"], r["market_cap"]) for r in rows])
    con.commit()


# ---------- compare ----------

def compare_last_two(con, threshold: float) -> tuple[pd.DataFrame, pd.DataFrame]:
    hist = pd.read_sql("SELECT * FROM snapshots", con)
    runs = sorted(hist["run_at"].unique())
    latest = hist[hist["run_at"] == runs[-1]].copy()
    if len(runs) > 1:
        prev = hist[hist["run_at"] == runs[-2]][["coin_id", "price"]].rename(columns={"price": "prev_price"})
        latest = latest.merge(prev, on="coin_id", how="left")
    else:
        latest["prev_price"] = float("nan")
    latest["change"] = latest["price"] - latest["prev_price"]
    latest["change_pct"] = (latest["change"] / latest["prev_price"] * 100).round(3)
    latest = latest.sort_values("market_cap", ascending=False)
    alerts = latest[latest["change_pct"].abs() >= threshold].copy()
    return latest, alerts


# ---------- report ----------

def style_header(ws):
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="1F6F43")
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 10), 40)


def write_report(con, latest: pd.DataFrame, alerts: pd.DataFrame, threshold: float, out: Path):
    out.parent.mkdir(parents=True, exist_ok=True)
    cols = {"symbol": "Symbol", "name": "Name", "price": "Price (USD)", "prev_price": "Previous (USD)",
            "change": "Change (USD)", "change_pct": "Change %"}
    latest_out = latest[list(cols)].rename(columns=cols)
    alerts_out = alerts[list(cols)].rename(columns=cols)

    hist = pd.read_sql("SELECT run_at, symbol, price FROM snapshots", con)
    top5 = latest["symbol"].head(5).tolist()
    pivot = (hist[hist["symbol"].isin(top5)]
             .pivot_table(index="run_at", columns="symbol", values="price")
             .reindex(columns=top5).reset_index().rename(columns={"run_at": "Run time"}))
    # index each coin to 100 at the first run so different price levels share one chart
    indexed = pivot.copy()
    for s in top5:
        indexed[s] = (pivot[s] / pivot[s].iloc[0] * 100).round(3)

    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        latest_out.to_excel(xw, sheet_name="Latest", index=False)
        alerts_out.to_excel(xw, sheet_name="Alerts", index=False)
        pivot.to_excel(xw, sheet_name="History", index=False)
        indexed.to_excel(xw, sheet_name="Chart data", index=False)

        ws = xw.sheets["Latest"]
        style_header(ws)
        for row in ws.iter_rows(min_row=2):
            pct = row[5].value
            if isinstance(pct, (int, float)) and pct == pct and pct != 0:
                fill = PatternFill("solid", fgColor=GREEN if pct > 0 else RED)
                for c in row[4:6]:
                    c.fill = fill
            for c in row[2:5]:
                c.number_format = '#,##0.00####'
            row[5].number_format = '+0.000"%";-0.000"%";0"%"'

        wa = xw.sheets["Alerts"]
        style_header(wa)
        wa.cell(row=wa.max_row + 2, column=1, value=f"Moves of {threshold}% or more since the previous run")

        for name in ("History", "Chart data"):
            style_header(xw.sheets[name])

        wc = xw.sheets["Chart data"]
        if wc.max_row > 2:
            chart = LineChart()
            chart.title = "Top 5 by market cap: price indexed to 100 at first run"
            chart.y_axis.title = "Index"
            chart.x_axis.title = "Run"
            chart.height, chart.width = 9, 20
            data = Reference(wc, min_col=2, max_col=1 + len(top5), min_row=1, max_row=wc.max_row)
            cats = Reference(wc, min_col=1, min_row=2, max_row=wc.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            xw.sheets["History"].add_chart(chart, f"{get_column_letter(len(top5) + 3)}2")


# ---------- main ----------

def run_once(args, run_no: int):
    run_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    rows = fetch_prices(args.coins)
    con = connect(args.db)
    save_snapshot(con, rows, run_at)
    latest, alerts = compare_last_two(con, args.threshold)
    write_report(con, latest, alerts, args.threshold, args.output)
    n_runs = con.execute("SELECT COUNT(DISTINCT run_at) FROM snapshots").fetchone()[0]
    con.close()

    if latest["prev_price"].isna().all():
        print(f"Run {run_no} at {run_at}: {len(rows)} prices saved (first run, nothing to compare yet)")
    else:
        moves = ", ".join(f"{r.symbol} {r.change_pct:+.2f}%" for r in alerts.itertuples()) or "none"
        print(f"Run {run_no} at {run_at}: {len(rows)} prices, {len(alerts)} moved "
              f">= {args.threshold}% ({moves})")
    print(f"  report updated: {args.output.name}  (history: {n_runs} runs)")


def main():
    ap = argparse.ArgumentParser(description="Track prices over time and flag big moves.")
    ap.add_argument("--coins", type=int, default=20, help="how many top coins to track")
    ap.add_argument("--threshold", type=float, default=0.5, help="alert threshold in percent")
    ap.add_argument("--every", type=float, default=None, help="minutes between runs (built-in schedule)")
    ap.add_argument("--runs", type=int, default=1, help="number of runs when using --every")
    ap.add_argument("--db", type=Path, default=HERE / "output" / "history.db")
    ap.add_argument("--output", type=Path, default=HERE / "output" / "price_report.xlsx")
    args = ap.parse_args()
    args.db.parent.mkdir(parents=True, exist_ok=True)

    total = args.runs if args.every else 1
    for i in range(1, total + 1):
        run_once(args, i)
        if i < total:
            time.sleep(args.every * 60)


if __name__ == "__main__":
    main()
